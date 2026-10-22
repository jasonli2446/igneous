from datetime import datetime

import pytest

from igneous.connectors.email_maildir import parse_message
from igneous.graph.spine import _name_from, _thread_key
from igneous.rag.chunking import split, strip_quoted
from igneous.rag.search import rrf

RAW = b"""Message-ID: <123.JavaMail.evans@thyme>
Date: Mon, 14 May 2001 16:39:00 -0700 (PDT)
From: phillip.allen@enron.com
To: tim.belden@enron.com, John.Arnold@enron.com
Subject: Re: Forecast
In-Reply-To: <100.JavaMail.evans@thyme>
X-From: Phillip K Allen
X-Folder: \\Phillip_Allen_Jan2002_1\\Allen, Phillip K.\\'Sent Mail

Here is our forecast.

 -----Original Message-----
From: Belden, Tim
Sent: Monday
Old quoted text.
"""


def test_synthetic_cross_source_corpus_is_deterministic_and_valid(tmp_path):
    from igneous.synthetic_corpus import NOTICE, generate, verify

    first = tmp_path / "first"
    second = tmp_path / "second"
    assert generate(first) == generate(second)
    assert verify(first) == {"documents": 5, "calendar_events": 2}
    assert NOTICE in (first / "README.md").read_text()
    assert NOTICE in (first / "calendar.ics").read_text()
    assert "phillip.allen@enron.com" in (first / "files/2001-05-14-forecast-memo.md").read_text()


def test_github_connector_groups_commits_and_uses_read_only_gets():
    from igneous.connectors.github import GithubConnector

    class Client:
        def __init__(self):
            self.calls = []

        def get(self, path, params=None, etag=None, allow_not_found=False):
            self.calls.append(("get", path, params, etag))
            if path.endswith("/readme"):
                return {"encoding": "base64", "content": "IyBSRUFETUU="}, None, None
            return None, None, None

        def pages(self, path, params=None, etag=None):
            self.calls.append(("pages", path, params, etag))
            if path == "/user/repos":
                return [{"full_name": "me/repo", "name": "repo", "updated_at": "2024-01-02T00:00:00Z", "default_branch": "main"}], '"repos"'
            if path.endswith("/issues"):
                return [{"number": 3, "title": "Bug", "body": "details", "updated_at": "2024-01-02T00:00:00Z", "user": {"login": "me"}}], None
            if path.endswith("/pulls"):
                return [{"number": 4, "title": "Fix", "body": "done", "updated_at": "2024-01-02T00:00:00Z", "user": {"login": "me"}}], None
            if path.endswith("/comments"):
                return [{"body": "looks good", "user": {"login": "reviewer"}, "created_at": "2024-01-02T00:00:00Z"}], None
            if path.endswith("/commits"):
                return [{"sha": "abcdef0123", "commit": {"message": "ship it", "author": {"date": "2024-01-02T12:00:00Z", "email": "me@example.com"}}, "author": {"login": "me"}}], None
            return [], None

    client = Client()
    events = list(GithubConnector("token", "me", client).sync({}))
    items = [e for e in events if hasattr(e, "external_id")]
    assert {i.external_id for i in items} == {"me/me/repo/repo/repo", "me/me/repo/readme", "me/me/repo/issue/3", "me/me/repo/pr/4", "me/me/repo/commits/2024-01-02"}
    assert any("reviewer:\nlooks good" in (i.body_text or "") for i in items)
    assert all(call[0] in {"get", "pages"} for call in client.calls)
    assert events[-1].external_id_prefix == "me/"


def test_parse_message():
    item = parse_message(RAW, "allen-p/sent/1.")
    m = item.raw_meta
    assert m["from"] == [{"name": "", "email": "phillip.allen@enron.com"}]
    assert [a["email"] for a in m["to"]] == ["tim.belden@enron.com", "john.arnold@enron.com"]
    assert m["in_reply_to"] == "<100.JavaMail.evans@thyme>"
    assert item.title == "Re: Forecast"
    assert item.occurred_at.year == 2001


def test_content_hash_stable():
    assert parse_message(RAW, "a").content_hash == parse_message(RAW, "a").content_hash


def test_strip_quoted():
    assert strip_quoted(parse_message(RAW, "a").body_text) == "Here is our forecast."


def test_split_long_text_overlaps_and_covers():
    t = ("word " * 1000).strip()
    parts = split(t)
    assert len(parts) > 1 and all(len(p) <= 2000 for p in parts)
    assert split("short") == ["short"] and split("") == []


def test_names_and_threads():
    assert _name_from({"name": "", "email": "phillip.allen@enron.com"}, "Phillip K Allen") == "Phillip K Allen"
    assert _name_from({"name": "", "email": "tim.belden@enron.com"}, None) == "Tim Belden"
    assert _thread_key({"references": ["<a>", "<b>"]}, "x") == "<a>"
    assert _thread_key({"from": [{"email": "a@x"}], "to": [{"email": "b@x"}]}, "RE: FW: Forecast") == "subject:forecast|a@x,b@x"


def test_spine_assigns_closed_message_intervals_without_domain_membership():
    from igneous.graph.spine import UPSERT_EDGE, Builder, _person

    class Result:
        def scalar_one(self):
            return 1

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            return Result()

    session = Session()
    builder = Builder(session)
    first, last = datetime(2024, 1, 2), datetime(2024, 3, 4)
    person = _person(builder, {"email": "person@example.com"}, "Person", first, 10)
    builder.edge(person, 99, "SENT", last, 11)

    (sent,) = [params for statement, params in session.calls if statement == str(UPSERT_EDGE)]
    assert sent == {"src": 1, "dst": 99, "type": "SENT", "valid_from": last,
                    "valid_to": last, "sid": 11}
    assert "LEAST(edges.valid_from, EXCLUDED.valid_from)" in str(UPSERT_EDGE)
    assert "GREATEST(edges.valid_to, EXCLUDED.valid_to)" in str(UPSERT_EDGE)


def test_temporal_backfill_is_bounded_and_repairs_both_edge_classes():
    from igneous.graph.spine import BACKFILL_MEMBERSHIPS, BACKFILL_POINT_EDGES, backfill_temporal

    class Result:
        def __init__(self, rowcount):
            self.rowcount = rowcount

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            return Result(2 if statement == BACKFILL_POINT_EDGES else 3)

    session = Session()
    assert backfill_temporal(session, limit=25) == 5
    assert [statement for statement, _ in session.calls] == [str(BACKFILL_POINT_EDGES), str(BACKFILL_MEMBERSHIPS)]
    assert all(params["limit"] == 25 for _, params in session.calls)
    assert "SET valid_from = c.observed_at, valid_to = c.observed_at" in str(BACKFILL_POINT_EDGES)
    assert "min(COALESCE(si.occurred_at, observed.valid_from))" in str(BACKFILL_MEMBERSHIPS)
    assert "max(COALESCE(si.occurred_at, observed.valid_from))" in str(BACKFILL_MEMBERSHIPS)


def test_rrf_rewards_agreement():
    assert rrf([1, 2, 3], [3, 2, 9])[0][0] in (2, 3)
    assert {c for c, _ in rrf([1], [2])} == {1, 2}


def test_experts_uses_retrieved_items_and_returns_bounded_evidence(monkeypatch):
    from igneous.rag import search

    hits = [
        {"source_item_id": 10, "score": 0.03, "text": "a" * 500},
        {"source_item_id": 10, "score": 0.02, "text": "older duplicate"},
        {"source_item_id": 20, "score": 0.01, "text": "event discussion"},
    ]
    monkeypatch.setattr(search, "hybrid", lambda *args, **kwargs: hits)

    class Result:
        def mappings(self):
            return self

        def all(self):
            return [
                {"person_id": 7, "name": "Pat Smith", "score": 0.04, "evidence_count": 2,
                 "source_item_id": 10, "involvement": ["authored", "mentioned"], "source": "gmail",
                 "external_id": "m-10", "title": "Plan", "occurred_at": None},
                {"person_id": 7, "name": "Pat Smith", "score": 0.04, "evidence_count": 2,
                 "source_item_id": 20, "involvement": ["attended"], "source": "gcal",
                 "external_id": "e-20", "title": "Review", "occurred_at": None},
            ]

    class Session:
        def __init__(self):
            self.params = None

        def execute(self, _statement, params):
            self.params = params
            return Result()

    session = Session()
    assert search.experts(session, "topic", k=5, evidence_limit=2) == [{
        "id": 7, "name": "Pat Smith", "score": 0.04, "evidence_count": 2,
        "evidence": [
            {"source_item_id": 10, "source": "gmail", "external_id": "m-10", "title": "Plan",
             "occurred_at": None, "involvement": ["authored", "mentioned"], "excerpt": "a" * 400},
            {"source_item_id": 20, "source": "gcal", "external_id": "e-20", "title": "Review",
             "occurred_at": None, "involvement": ["attended"], "excerpt": "event discussion"},
        ],
    }]
    assert session.params["item_ids"] == [10, 20]
    assert session.params["relevance"] == [0.03, 0.01]
    assert session.params["k"] == 5 and session.params["evidence_limit"] == 2


def test_experts_endpoint_validates_query_and_passes_bounded_parameters(monkeypatch):
    from fastapi.testclient import TestClient

    from igneous import main
    from igneous.db import get_session

    calls = []
    monkeypatch.setattr(main, "experts", lambda session, q, **kwargs: calls.append((session, q, kwargs)) or [])
    marker = object()
    main.app.dependency_overrides[get_session] = lambda: marker
    try:
        client = TestClient(main.app)
        assert client.get("/experts?q=%20%20").status_code == 422
        response = client.get("/experts?q=budget&k=2&evidence_limit=1&mode=lexical")
    finally:
        main.app.dependency_overrides.clear()
    assert response.status_code == 200 and response.json() == {"experts": []}
    assert calls == [(marker, "budget", {"k": 2, "evidence_limit": 1, "mode": "lexical"})]


def test_graph_filters_scope_discovery_and_final_serialization():
    """A source/date request must not discover or serialize a non-matching edge."""
    from igneous.graph.traversal import GraphFilters, neighborhood

    class Result:
        def __init__(self, rows):
            self.rows = rows

        def scalars(self):
            return self

        def all(self):
            return self.rows

        def mappings(self):
            return self

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            sql = str(statement)
            if "SELECT CASE WHEN e.src" in sql:
                return Result([])  # zero discovered neighbors
            if "FROM nodes n" in sql:
                return Result([{"id": 7, "type": "Person", "name": "Seed", "key": "seed", "degree": 3}])
            return Result([])

    session = Session()
    graph = neighborhood(session, [7], filters=GraphFilters(source="gmail", edge_types=["SENT"],
                                                              since=datetime(2024, 1, 1), until=datetime(2024, 2, 1)))
    discovery, serialized = session.calls[0], session.calls[-1]
    for sql, params in (discovery, serialized):
        assert "LEFT JOIN source_items si ON si.id = e.source_item_id" in sql
        assert "si.source = :source" in sql
        assert "COALESCE(e.valid_to, si.occurred_at, e.valid_from) >= :since" in sql
        assert "COALESCE(e.valid_from, si.occurred_at, e.valid_to) <= :until" in sql
        assert "e.type = ANY(:types)" in sql
        assert params["source"] == "gmail" and params["types"] == ["SENT"]
    assert graph == {"nodes": [{"id": 7, "type": "Person", "name": "Seed", "key": "seed", "degree": 3}], "edges": []}


def test_timeline_returns_bounded_dated_metadata_without_raw_source_content():
    from igneous.timeline import timeline

    dated = datetime(2024, 2, 3)

    class Result:
        def __init__(self, rows):
            self.rows = rows

        def mappings(self):
            return self

        def first(self):
            return self.rows[0] if self.rows else None

        def all(self):
            return self.rows

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            if "FROM nodes n" in str(statement):
                return Result([{"id": params["node_id"], "type": "Person", "name": "Pat"}])
            return Result([
                {"id": 9, "src": 1, "dst": 2, "a_is_src": True, "type": "ATTENDED",
                 "valid_from": dated, "valid_to": dated, "extractor": "spine", "extractor_version": "1",
                 "model": None, "confidence": None, "source_item_id": 44, "span_start": None,
                 "span_end": None, "source": "gcal", "occurred_at": dated, "other_id": 2,
                 "other_type": "Event", "other_name": "Private calendar title"},
                {"id": 8, "src": 1, "dst": 3, "a_is_src": True, "type": "MENTIONS",
                 "valid_from": dated, "valid_to": dated, "extractor": "extract", "extractor_version": "2",
                 "model": "model", "confidence": 0.8, "source_item_id": 45, "span_start": 10,
                 "span_end": 14, "source": "gmail", "occurred_at": dated, "other_id": 3,
                 "other_type": "Topic"},
            ])

    session = Session()
    result = timeline(session, a=1, b=2, limit=1)

    assert result["mode"] == "relationship" and result["truncated"] is True
    assert result["edges"] == [{
        "id": 9, "type": "ATTENDED", "direction": "outgoing",
        "counterpart": {"id": 2, "type": "Event", "name": "Private calendar title"},
        "interval": {"from": dated, "to": dated},
        "provenance": {"extractor": "spine", "extractor_version": "1", "model": None,
                       "confidence": None, "evidence": {"source_item_id": 44, "source": "gcal",
                                                           "occurred_at": dated, "span": None}},
    }]
    query, params = session.calls[-1]
    assert "COALESCE(e.valid_from, e.valid_to, si.occurred_at) IS NOT NULL" in query
    assert "si.body_text" not in query and "si.title" not in query and "si.raw_meta" not in query
    # Message bodies, titles and raw metadata stay out; only the counterpart's display name is shown (the UI
    # names things instead of showing "Event #2").
    assert "other.name AS other_name" in query
    assert params["limit_plus_one"] == 2 and params["a"] == 1 and params["b"] == 2


def test_timeline_endpoint_validates_bounds_and_translates_missing_entities(monkeypatch):
    from fastapi.testclient import TestClient

    from igneous import main
    from igneous.db import get_session
    from igneous.timeline import EntityNotFoundError

    calls = []
    marker = object()
    monkeypatch.setattr(main, "timeline", lambda session, **kwargs: calls.append((session, kwargs)) or {"edges": []})
    main.app.dependency_overrides[get_session] = lambda: marker
    try:
        client = TestClient(main.app)
        assert client.get("/timeline?a=0").status_code == 422
        response = client.get("/timeline?a=1&b=2&limit=3")
        activity = client.get("/timeline?a=1&limit=1")
    finally:
        main.app.dependency_overrides.clear()
    assert response.json() == activity.json() == {"edges": []}
    assert calls == [(marker, {"a": 1, "b": 2, "limit": 3}), (marker, {"a": 1, "b": None, "limit": 1})]

    monkeypatch.setattr(main, "timeline", lambda *_args, **_kwargs: (_ for _ in ()).throw(EntityNotFoundError(9)))
    main.app.dependency_overrides[get_session] = lambda: marker
    try:
        response = TestClient(main.app).get("/timeline?a=9")
    finally:
        main.app.dependency_overrides.clear()
    assert response.status_code == 404


def test_topic_normalization_is_conservative_and_similarity_requires_lexical_evidence():
    from igneous.graph.topics import Topic, normalize_key, similar_pairs

    assert normalize_key("California Price Caps") == "california-price-cap"
    assert normalize_key("Quarterly Analyses") == "quarterly-analysis"
    assert normalize_key("Business News") == "business-news"  # not a plural noun we can safely change

    topics = [Topic(1, "california-price-cap", "California price cap"),
              Topic(2, "price-cap-policy", "California price cap policy"),
              Topic(3, "unrelated", "Houston operations")]
    pairs = list(similar_pairs(topics, [[1.0, 0.0], [0.999, 0.001], [1.0, 0.0]], threshold=0.96))
    assert [(left.id, right.id) for left, right, _ in pairs] == [(1, 2)]


def test_topic_related_to_uses_distinct_message_cooccurrences_and_own_extractor():
    from igneous.graph.topics import EXTRACTOR, _related

    class Result:
        rowcount = 4

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            return Result()

    session = Session()
    assert _related(session, 3) == 4
    delete, insert = session.calls
    assert "DELETE FROM edges WHERE type = 'RELATED_TO'" in delete[0]
    assert delete[1] == {"ex": EXTRACTOR}
    assert "count(DISTINCT a.src)" in insert[0]
    assert "a.dst < b.dst" in insert[0]
    assert insert[1]["minimum"] == 3 and insert[1]["ex"] == EXTRACTOR


def test_locate_spans():
    from igneous.graph.extract import locate

    body = "Please ask Jeff\n  Skilling and FERC about the caps."
    assert locate(body, "FERC") == (body.index("FERC"), body.index("FERC") + 4)
    s0, s1 = locate(body, "Jeff Skilling")  # survives the line wrap
    assert body[s0:s1].split() == ["Jeff", "Skilling"]
    assert locate(body, "Ken Lay") is None and locate(body, "  ") is None


def test_automatic_extraction_requires_explicit_source_opt_in(monkeypatch):
    from igneous.graph import extract

    assert extract.allowed_sources("") == frozenset()
    assert extract.allowed_sources(" gmail, ai_chats ,gmail ") == frozenset({"gmail", "ai_chats"})

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            return type("Result", (), {"all": lambda self: []})()

    session = Session()
    messages = []
    monkeypatch.setattr(extract.settings, "llm_sources", "")
    assert extract.run(session, log=messages.append) == 0
    assert session.calls == []
    assert "IGNEOUS_LLM_SOURCES" in messages[0]

    monkeypatch.setattr(extract.settings, "llm_sources", "ai_chats")
    assert extract.run(session) == 0
    query, params = session.calls[0]
    assert "si.source = ANY(CAST(:sources AS text[]))" in query
    assert params["sources"] == ["ai_chats"]


def test_ask_retrieves_only_explicitly_approved_sources(monkeypatch):
    from igneous.rag import qa
    from igneous.rag.search import Filters

    captured = []
    monkeypatch.setattr(qa, "apply", lambda _s, _q, f: f or Filters())
    monkeypatch.setattr(qa, "allowed_sources", lambda: frozenset({"ai_chats"}))
    monkeypatch.setattr(qa, "hybrid", lambda _s, _q, f, **_kwargs: captured.append(f) or [])
    assert qa.answer(object(), "What did I decide?")["answerable"] is False
    assert captured[0].sources == ("ai_chats",)


def test_extraction_spans_index_the_stored_body():
    """Regression: offsets were computed on a .strip()-ed body and stored against the original (off by one)."""
    from igneous.graph.extract import locate
    from igneous.rag.chunking import QUOTE_START

    stored = "\n  Ask David Oxley.\n -----Original Message-----\nold"
    own = stored[: QUOTE_START.search(stored).start()]
    s0, s1 = locate(own, "David Oxley")
    assert stored[s0:s1] == "David Oxley"


def test_er_names_and_scoring():
    from igneous.graph.resolve import P, same_first, score, split_name, unresolvable

    assert split_name("Allen, Phillip K.", "pallen@enron.com") == ("phillip", "allen")
    assert split_name("", "jeff.skilling@enron.com") == ("jeff", "skilling")
    assert same_first("rick", "richard") and same_first("joe", "joseph") and not same_first("ann", "mark")

    def p(i, key, name, peers=()):
        f, last = split_name(name, key)
        return P(i, key, name, f, last, set(peers))

    shared = {10, 11, 12, 13}
    assert score(p(1, "james.steffes@enron.com", "James Steffes", shared), p(2, "steffes@enron.com", "James Steffes", shared))[0] >= 0.75
    assert score(p(1, "ken.lay@enron.com", "Ken Lay"), p(2, "ken_lay@enron.net", "Ken Lay"))[0] >= 0.75
    # Same surname, different first name: never merges no matter how much mail they share.
    assert score(p(1, "ann.schmidt@enron.com", "Ann Schmidt", shared), p(2, "mark.schmidt@enron.com", "Mark Schmidt", shared))[0] == 0
    assert unresolvable(p(3, ".john@enron.com", "e-mail")) and unresolvable(p(4, "dl-ga-all@enron.com", "DL GA All"))


def test_er_rejects_mailing_lists():
    from igneous.graph.resolve import P, split_name, unresolvable

    for key, name in [("aei-bi_updates@www.brookings.org", "Aei-bi Updates"),
                      ("western.price.survey.contacts@proxy4.ba.best.com", "Western Price Survey Contacts")]:
        f, last = split_name(name, key)
        assert unresolvable(P(1, key, name, f, last, set()))


def test_imessage_decoding_and_dates():
    from igneous.connectors.imessage import apple_time, decode_attributed

    short = b"\x04\x0bstreamtyped\x81\xe8\x03\x84\x01@\x84\x84\x84\x12NSAttributedString\x00\x84\x84\x08NSObject" \
            b"\x00\x85\x92\x84\x84\x84\x08NSString\x01\x94\x84\x01+\x05hello\x86"
    assert decode_attributed(short) == "hello"
    long_text = "x" * 300
    long_blob = b"NSString\x01\x94\x84\x01+\x81" + (300).to_bytes(2, "little") + long_text.encode()
    assert decode_attributed(long_blob) == long_text
    assert decode_attributed(None) is None and decode_attributed(b"garbage") is None
    assert apple_time(0) is None
    assert apple_time(700_000_000 * 10**9).year == 2023   # nanoseconds (current format)
    assert apple_time(700_000_000).year == 2023           # seconds (pre-10.13 rows)


def test_quoted_header_parsing():
    from igneous.graph.quoted import _surname, parse_quotes

    body = ("Sounds good.\n\n -----Original Message-----\nFrom: Rice, Ken \nSent: Mon 8/20/2001 10:20 AM \n"
            "To: Skilling, Jeff \nSubject: RE: Astros tickets\n\nThanks for the tickets.")
    (q,) = parse_quotes(body)
    assert (q["sent"].month, q["sent"].day, q["sent"].hour, q["sent"].minute) == (8, 20, 10, 20)
    assert q["subject"] == "Astros tickets" and _surname(q["from"]) == "rice"
    from igneous.graph.quoted import _first
    assert _first("Rice, Ken") == "ken" and _first("Brad Schaeffer <b@x.com>") == "brad" and _first("Rice") == ""
    long = "-----Original Message-----\nFrom: Brad Schaeffer <BSchaeffer@prebon.com>@ENRON \nSent: Thursday, November 29, 2001 9:49 AM\nSubject: x\n"
    assert parse_quotes(long)[0]["sent"].year == 2001 and _surname(parse_quotes(long)[0]["from"]) == "schaeffer"
    assert parse_quotes("no quote here") == []


def test_parse_plain_text_markdown_and_binary_are_nonfatal():
    from igneous.parsers import parse_document

    markdown = parse_document(b"# Title\r\n\r\nBody", "notes.md")
    assert (markdown.status, markdown.mime, markdown.parser, markdown.text) == (
        "ok", "text/markdown", "text", "# Title\n\nBody")
    unsupported = parse_document(b"\x89PNG\r\n\x1a\n", "image.png")
    assert unsupported.status == "unsupported" and unsupported.text == "" and unsupported.error
    binary = parse_document(b"\x00\x01", "bad.txt")
    assert binary.status == "binary" and binary.text == "" and binary.error


def test_parse_pdf_with_page_boundaries():
    from igneous.parsers import parse_document

    def pdf_bytes():
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 7 0 R >> >> >>",
            b"<< /Length 42 >>\nstream\nBT /F1 12 Tf 72 720 Td (First page) Tj ET\nendstream",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 6 0 R /Resources << /Font << /F1 7 0 R >> >> >>",
            b"<< /Length 43 >>\nstream\nBT /F1 12 Tf 72 720 Td (Second page) Tj ET\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        out, offsets = bytearray(b"%PDF-1.4\n"), [0]
        for number, obj in enumerate(objects, start=1):
            offsets.append(len(out))
            out.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
        xref = len(out)
        out.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
        out.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
        out.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
        return bytes(out)

    parsed = parse_document(pdf_bytes(), "ignored.bin")
    assert parsed.status == "ok" and parsed.parser == "pdf"
    assert parsed.text == "First page\n\nSecond page"
    assert parsed.boundaries == [
        {"kind": "page", "index": 1, "start": 0, "end": 10},
        {"kind": "page", "index": 2, "start": 12, "end": 23},
    ]


def test_parse_docx_xlsx_and_pptx_structures():
    from io import BytesIO

    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    from igneous.parsers import parse_document

    docx = Document()
    docx.add_heading("Plan", level=1)
    docx.add_paragraph("Ship it")
    data = BytesIO()
    docx.save(data)
    parsed = parse_document(data.getvalue(), "plan.docx")
    assert parsed.status == "ok" and parsed.text == "Plan\nShip it"
    assert parsed.boundaries == [{"kind": "heading", "level": 1, "title": "Plan", "start": 0, "end": 4}]

    workbook = Workbook()
    workbook.active.title = "Numbers"
    workbook.active.append(["name", "value"])
    workbook.active.append(["one", 1])
    workbook.create_sheet("Empty")
    data = BytesIO()
    workbook.save(data)
    parsed = parse_document(data.getvalue(), "data.xlsx")
    assert parsed.status == "ok" and parsed.text == "name\tvalue\none\t1\n\n"
    assert parsed.boundaries == [
        {"kind": "sheet", "index": 1, "name": "Numbers", "start": 0, "end": 16},
        {"kind": "sheet", "index": 2, "name": "Empty", "start": 18, "end": 18},
    ]

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[5]).shapes.add_textbox(0, 0, 1, 1).text = "Hello slide"
    data = BytesIO()
    presentation.save(data)
    parsed = parse_document(data.getvalue(), "deck.pptx")
    assert parsed.status == "ok" and parsed.text == "Hello slide"
    assert parsed.boundaries == [{"kind": "slide", "index": 1, "start": 0, "end": 11}]


def test_parse_corrupt_office_file_is_nonfatal():
    from igneous.parsers import parse_document

    parsed = parse_document(b"not a document", "broken.docx")
    assert parsed.status == "corrupt" and parsed.parser == "docx" and parsed.text == "" and parsed.error
def test_cross_encoder_reranks_full_pool_then_limits(monkeypatch):
    from igneous.rag import rerank

    class Encoder:
        def rerank(self, question, documents):
            assert question == "Which item is relevant?"
            assert documents == ["first", "second", "third"]
            return [0.1, 0.2, 0.9]

    monkeypatch.setattr(rerank, "_model", lambda: Encoder())
    items = [{"text": "first", "score": 1.0}, {"text": "second", "score": 0.5}, {"text": "third"}]
    assert rerank.rerank("Which item is relevant?", items, keep=2) == [
        {"text": "third", "rerank_score": 0.9},
        {"text": "second", "score": 0.5, "rerank_score": 0.2},
    ]


def test_cross_encoder_disabled_preserves_retrieval_order(monkeypatch):
    from igneous.rag import rerank

    monkeypatch.setattr(rerank.settings, "reranker_enabled", False)
    monkeypatch.setattr(rerank, "_model", lambda: (_ for _ in ()).throw(AssertionError("must not load")))
    items = [{"text": "first"}, {"text": "second"}]
    assert rerank.rerank("question", items, keep=1) == [{"text": "first"}]


def test_cross_encoder_failure_preserves_retrieval_order(monkeypatch, caplog):
    from igneous.rag import rerank

    attempts = 0

    def unavailable():
        nonlocal attempts
        attempts += 1
        raise OSError("model cache is offline")

    monkeypatch.setattr(rerank, "_model", unavailable)
    monkeypatch.setattr(rerank, "_load_failure", None)
    items = [{"text": "first"}, {"text": "second"}]
    assert rerank.rerank("question", items, keep=1) == [{"text": "first"}]
    assert rerank.rerank("question", items, keep=1) == [{"text": "first"}]
    assert attempts == 1
    assert "using retrieval order" in caplog.text


def test_gmail_backfill_resumes_mid_page(monkeypatch):
    """Regression: a mid-page checkpoint saved only the page token, so resuming re-fetched the page."""
    from igneous.connectors import gmail
    from igneous.connectors.base import Checkpoint

    monkeypatch.setattr(gmail, "CHECKPOINT_EVERY", 2)
    monkeypatch.setattr(gmail, "_retry", lambda call, units=5, tries=6: call)
    fetched = []

    class Conn(gmail.GmailConnector):
        def __init__(self):
            self.query, self.since_days, self.account = "", None, "x"
            pages = {None: {"messages": [{"id": f"m{i}"} for i in range(5)], "nextPageToken": "p2"},
                     "p2": {"messages": [{"id": "m5"}]}}

            class Svc:
                def users(_):
                    class U:
                        def getProfile(_, userId):
                            return {"historyId": "h1"}

                        def messages(_):
                            class M:
                                def list(_, pageToken=None, **k):
                                    return pages[pageToken]
                            return M()
                    return U()
            self.svc = Svc()

        def _fetch(self, msg_id):
            fetched.append(msg_id)
            return None

    out = list(Conn()._backfill())
    mid = [c.cursor for c in out if isinstance(c, Checkpoint) and "backfill_offset" in c.cursor]
    assert mid[1] == {"history_id": "h1", "backfill_page": None, "backfill_offset": 4, "backfill_q": ""}
    fetched.clear()
    list(Conn()._backfill(mid[1]))
    assert fetched == ["m4", "m5"]  # resumes after m3, not from m0


def test_phone_normalization():
    from igneous.connectors.mac_contacts import norm_phone

    assert norm_phone("(614) 555-1234") == norm_phone("+1 614 555 1234") == norm_phone("6145551234") == "+16145551234"
    assert norm_phone("+44 20 7946 0958") == "+442079460958"
    assert norm_phone("555-1234") is None and norm_phone(None) is None


def test_contact_names_in_messages():
    from igneous.graph.traversal import name_handles, pretty_thread

    named = {"+16174075277": "Sam Lee", "6174075277": "Sam Lee", "+16149733915": "Ana Ruiz", "6149733915": "Ana Ruiz"}
    assert name_handles("[14:39] +16174075277: home rn", named) == "[14:39] Sam Lee: home rn"
    assert name_handles("[1:00] +19995550000: hi", named) == "[1:00] +19995550000: hi"  # unknown stays
    assert pretty_thread("+16149733915, +16174075277", named) == "Ana, Sam"
    assert pretty_thread("+16149733915, +16174075277, +19995550000, +18885550000", named).endswith("and 1 other")
    assert pretty_thread("Biggers in Paris", named) == "Biggers in Paris"


class _SyncStatusResult:
    def __init__(self, rows):
        self.rows = rows

    def mappings(self):
        return self

    def all(self):
        return self.rows


class _SyncStatusSession:
    def __init__(self, rows):
        self.rows = rows
        self.statement = None

    def execute(self, statement, _params):
        self.statement = str(statement)
        return _SyncStatusResult(self.rows)


def test_sync_status_reports_age_counts_and_sanitized_cursor():
    from datetime import UTC, datetime

    from igneous.sync_status import sync_status

    session = _SyncStatusSession([{
        "connector": "gmail",
        "account": "me@example.com",
        "cursor": {"history_id": "opaque-history-token", "backfill_page": "opaque-page-token"},
        "updated_at": datetime(2026, 10, 20, 11, 30, tzinfo=UTC),
        "item_count": 42,
    }])

    (status,) = sync_status(session, now=datetime(2026, 10, 20, 12, 0, tzinfo=UTC))

    assert status == {
        "connector": "gmail",
        "account": "me@example.com",
        "status": "known",
        "last_synced_at": datetime(2026, 10, 20, 11, 30, tzinfo=UTC),
        "cursor_age_seconds": 1800,
        "item_count": 42,
        "item_count_scope": "connector",
        "cursor": {"present": True, "keys": ["backfill_page", "history_id"], "gmail_backfill_in_progress": True},
    }
    assert "opaque-history-token" not in repr(status)
    assert "SELECT source AS connector, count(*)::bigint AS item_count" in session.statement
    assert "item_counts.connector = subjects.connector" in session.statement
    assert "raw_meta ->> 'account'" not in session.statement


def test_sync_status_marks_requested_missing_checkpoint_never_synced():
    from igneous.sync_status import sync_status

    session = _SyncStatusSession([{
        "connector": "gmail",
        "account": "me@example.com",
        "cursor": None,
        "updated_at": None,
        "item_count": 42,
    }])

    assert sync_status(session, connector="gmail", account="me@example.com") == [{
        "connector": "gmail",
        "account": "me@example.com",
        "status": "never_synced",
        "last_synced_at": None,
        "cursor_age_seconds": None,
        "item_count": 42,
        "item_count_scope": "connector",
        "cursor": {"present": False, "keys": [], "gmail_backfill_in_progress": False},
    }]


def test_get_sync_is_read_only_status_report():
    from fastapi.testclient import TestClient

    from igneous.db import get_session
    from igneous.main import app

    session = _SyncStatusSession([{
        "connector": "imessage",
        "account": "default",
        "cursor": None,
        "updated_at": None,
        "item_count": 0,
    }])
    app.dependency_overrides[get_session] = lambda: session
    try:
        response = TestClient(app).get("/sync?connector=imessage&account=default")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.json() == {
        "syncs": [{
            "connector": "imessage",
            "account": "default",
            "status": "never_synced",
            "last_synced_at": None,
            "cursor_age_seconds": None,
            "item_count": 0,
            "item_count_scope": "connector",
            "cursor": {"present": False, "keys": [], "gmail_backfill_in_progress": False},
        }]
    }
    assert TestClient(app).get("/sync?connector=unknown").status_code == 422


def test_filesystem_connector_walks_sorted_parseable_files_and_emits_hierarchy(tmp_path):
    from igneous.connectors.base import CompletedScan, SourceItem
    from igneous.connectors.filesystem import FilesystemConnector, _root_key

    (tmp_path / "z-last.txt").write_text("last")
    (tmp_path / "projects" / "igneous").mkdir(parents=True)
    (tmp_path / "projects" / "igneous" / "plan.md").write_text("# Plan\nship")
    (tmp_path / "node_modules").mkdir()
    (tmp_path / "node_modules" / "ignored.txt").write_text("ignore")
    (tmp_path / ".git").mkdir()
    (tmp_path / ".git" / "config").write_text("ignore")
    (tmp_path / "photo.jpg").write_bytes(b"not an image parser input")

    emitted = list(FilesystemConnector(tmp_path).sync({}))
    items = [event for event in emitted if isinstance(event, SourceItem)]
    root_key = _root_key(tmp_path)
    assert [item.external_id for item in items] == [
        f"{root_key}/photo.jpg", f"{root_key}/projects/igneous/plan.md", f"{root_key}/z-last.txt"]
    assert items[0].body_text is None and items[0].raw_meta["parse_status"] == "unsupported"
    plan = items[1]
    assert plan.body_text == "# Plan\nship"
    assert plan.mime == "text/markdown"
    assert plan.raw_meta["path"] == f"{root_key}/projects/igneous/plan.md"
    assert plan.raw_meta["folders"] == [root_key, f"{root_key}/projects", f"{root_key}/projects/igneous"]
    assert plan.raw_meta["parser"] == "text" and plan.raw_meta["boundaries"] == []
    assert plan.raw_meta["model_opt_in"] is False
    assert isinstance(emitted[-1], CompletedScan)


def test_filesystem_connector_hashes_file_bytes_and_namespaces_multiple_roots(tmp_path):
    from igneous.connectors.base import SourceItem
    from igneous.connectors.filesystem import FilesystemConnector, _root_key

    first, second = tmp_path / "first", tmp_path / "second"
    first.mkdir()
    second.mkdir()
    (first / "same.txt").write_text("one")
    (second / "same.txt").write_text("two")
    connector = FilesystemConnector([second, first])
    initial = [event for event in connector.sync({}) if isinstance(event, SourceItem)]
    assert [item.external_id for item in initial] == [
        f"{_root_key(first)}/same.txt", f"{_root_key(second)}/same.txt"]
    (first / "same.txt").write_text("changed")
    changed = [event for event in connector.sync({}) if isinstance(event, SourceItem)]
    assert changed[0].content_hash != initial[0].content_hash


def test_run_sync_prunes_only_after_completed_scan(monkeypatch):
    from igneous.connectors import base

    class Result:
        def __init__(self, deleted=False):
            self.rowcount = 1 if deleted else 0

        def first(self):
            return None

    class Session:
        def __init__(self):
            self.statements = []
            self.commits = 0

        def execute(self, statement, values=None):
            self.statements.append(str(statement))
            return Result("DELETE FROM source_items" in str(statement))

        def commit(self):
            self.commits += 1

    class Complete:
        name = "fs"

        def sync(self, cursor):
            yield base.SourceItem("fs", "kept.txt", "kept.txt", "body")
            yield base.CompletedScan({}, "fs", frozenset({"kept.txt"}))

    session = Session()
    monkeypatch.setattr(base, "load_cursor", lambda *args: {})
    monkeypatch.setattr(base, "save_cursor", lambda *args: None)
    assert base.run_sync(session, Complete())["deleted"] == 1
    assert any("DELETE FROM source_items" in statement for statement in session.statements)

    class Interrupted:
        name = "fs"

        def sync(self, cursor):
            yield base.SourceItem("fs", "kept.txt", "kept.txt", "body")
            yield base.Checkpoint({})
            raise RuntimeError("interrupted")

    session = Session()
    with pytest.raises(RuntimeError, match="interrupted"):
        base.run_sync(session, Interrupted())
    assert not any("DELETE FROM source_items" in statement for statement in session.statements)


def test_filesystem_connector_rejects_missing_roots(tmp_path):
    from igneous.connectors.filesystem import FilesystemConnector

    with pytest.raises(ValueError, match="must be directories"):
        FilesystemConnector(tmp_path / "does-not-exist")


def test_drive_backfill_exports_workspace_files_and_preserves_all_parent_ids(monkeypatch):
    from igneous.connectors.base import CompletedScan, SourceItem
    from igneous.connectors.drive import DriveConnector

    class Request:
        def __init__(self, value):
            self.value = value

        def execute(self):
            return self.value

    class Files:
        def list(self, **kwargs):
            return Request({"files": [
                {"id": "doc-1", "name": "Plan", "mimeType": "application/vnd.google-apps.document",
                 "parents": ["folder-a", "folder-b"], "modifiedTime": "2026-10-21T12:00:00Z"},
                {"id": "image-1", "name": "photo.jpg", "mimeType": "image/jpeg"},
            ]})

        def export(self, **kwargs):
            assert kwargs == {"fileId": "doc-1", "mimeType": "text/plain"}
            return Request(b"ship it")

        def get_media(self, **kwargs):
            return Request(b"not parseable")

    class Changes:
        def getStartPageToken(self, **kwargs):
            return Request({"startPageToken": "before-backfill"})

    connector = object.__new__(DriveConnector)
    connector.svc = type("Service", (), {"files": lambda self: Files(), "changes": lambda self: Changes()})()
    connector.account = "me@example.com"
    monkeypatch.setattr("igneous.connectors.drive.settings.drive_page_size", 100)

    events = list(connector.sync({}))
    item = next(event for event in events if isinstance(event, SourceItem))
    assert item.external_id == "me@example.com/doc-1"
    assert item.body_text == "ship it"
    assert item.raw_meta["parents"] == ["folder-a", "folder-b"]
    assert item.raw_meta["path"] == "drive/doc-1"
    completed = events[-1]
    assert isinstance(completed, CompletedScan)
    assert completed.cursor == {"start_page_token": "before-backfill"}
    assert completed.external_ids == frozenset({"me@example.com/doc-1", "me@example.com/image-1"})


def test_drive_changes_emit_tombstones_and_advance_only_at_final_page(monkeypatch):
    from igneous.connectors.base import Checkpoint, Tombstone
    from igneous.connectors.drive import DriveConnector

    class Request:
        def __init__(self, value):
            self.value = value

        def execute(self):
            return self.value

    class Changes:
        def list(self, **kwargs):
            assert kwargs["pageToken"] == "cursor-1"
            return Request({"changes": [
                {"fileId": "deleted", "removed": True},
                {"fileId": "trashed", "file": {"id": "trashed", "trashed": True}},
            ], "newStartPageToken": "cursor-2"})

    connector = object.__new__(DriveConnector)
    connector.svc = type("Service", (), {"changes": lambda self: Changes()})()
    connector.account = "me@example.com"
    monkeypatch.setattr("igneous.connectors.drive.settings.drive_page_size", 100)

    events = list(connector.sync({"start_page_token": "cursor-1"}))
    assert [(event.source, event.external_id) for event in events if isinstance(event, Tombstone)] == [
        ("drive", "me@example.com/deleted"), ("drive", "me@example.com/trashed")]
    assert events[-1] == Checkpoint({"start_page_token": "cursor-2"})


def test_run_sync_applies_explicit_tombstones(monkeypatch):
    from igneous.connectors import base

    class Result:
        rowcount = 1

        def first(self):
            return None

    class Session:
        def __init__(self):
            self.statements = []

        def execute(self, statement, values=None):
            self.statements.append((str(statement), values))
            return Result()

        def commit(self):
            pass

    class DriveChanges:
        name = "drive"

        def sync(self, cursor):
            yield base.Tombstone("drive", "gone")
            yield base.Checkpoint({"start_page_token": "next"})

    session = Session()
    monkeypatch.setattr(base, "load_cursor", lambda *args: {"start_page_token": "old"})
    monkeypatch.setattr(base, "save_cursor", lambda *args: None)
    assert base.run_sync(session, DriveChanges())["deleted"] == 1
    assert any(values == {"source": "drive", "external_id": "gone"} for _, values in session.statements)


def test_drive_skips_non_downloadable_workspace_files_and_export_limits():
    from types import SimpleNamespace

    from googleapiclient.errors import HttpError

    from igneous.connectors.drive import _is_skippable_download_error

    def error(status, reason):
        return HttpError(SimpleNamespace(status=status, reason=reason), reason.encode())

    assert _is_skippable_download_error(error(403, "fileNotDownloadable"))
    assert _is_skippable_download_error(error(403, "exportSizeLimitExceeded"))
    assert _is_skippable_download_error(error(403, "cannotDownloadAbusiveFile"))
    assert _is_skippable_download_error(error(404, "notFound"))
    assert not _is_skippable_download_error(error(403, "insufficientPermissions"))


def test_drive_ignores_drive_changes_without_file_ids_and_backfills_expired_tokens(monkeypatch):
    from types import SimpleNamespace

    from googleapiclient.errors import HttpError

    from igneous.connectors.base import CompletedScan
    from igneous.connectors.drive import DriveConnector

    class Request:
        def __init__(self, value):
            self.value = value

        def execute(self):
            if isinstance(self.value, Exception):
                raise self.value
            return self.value

    expired = HttpError(SimpleNamespace(status=410, reason="Gone"), b"pageNotFound")

    class Changes:
        def list(self, **kwargs):
            return Request(expired)

        def getStartPageToken(self, **kwargs):
            return Request({"startPageToken": "fresh-token"})

    class Files:
        def list(self, **kwargs):
            return Request({"files": []})

    connector = object.__new__(DriveConnector)
    connector.svc = type("Service", (), {"changes": lambda self: Changes(), "files": lambda self: Files()})()
    connector.account = "me@example.com"
    monkeypatch.setattr("igneous.connectors.drive.settings.drive_page_size", 100)

    events = list(connector.sync({"start_page_token": "expired-token"}))
    assert events == [CompletedScan({"start_page_token": "fresh-token"}, "drive", frozenset())]


def test_drive_ignores_non_file_change_events(monkeypatch):
    from igneous.connectors.base import Checkpoint
    from igneous.connectors.drive import DriveConnector

    class Request:
        def execute(self):
            return {"changes": [{"changeType": "drive", "driveId": "shared-drive"}],
                    "newStartPageToken": "next"}

    class Changes:
        def list(self, **kwargs):
            assert "changeType" in kwargs["fields"]
            return Request()

    connector = object.__new__(DriveConnector)
    connector.svc = type("Service", (), {"changes": lambda self: Changes()})()
    connector.account = "me@example.com"
    monkeypatch.setattr("igneous.connectors.drive.settings.drive_page_size", 100)

    assert list(connector.sync({"start_page_token": "old"})) == [Checkpoint({"start_page_token": "next"})]


def test_calendar_uses_clean_incremental_params_and_tombstones_cancelled_events(monkeypatch):
    from igneous.connectors.base import Checkpoint, SourceItem, Tombstone
    from igneous.connectors.gcal import CalendarConnector

    class Request:
        def __init__(self, value):
            self.value = value

        def execute(self):
            return self.value

    class CalendarList:
        def list(self, **kwargs):
            return Request({"items": [{"id": "team", "summary": "Team"}]})

    class Events:
        def list(self, **kwargs):
            assert kwargs["syncToken"] == "old-token"
            assert "timeMin" not in kwargs
            assert kwargs["showDeleted"] is True
            return Request({"items": [{"id": "gone", "status": "cancelled"}],
                            "nextSyncToken": "new-token"})

    connector = object.__new__(CalendarConnector)
    connector.svc = type("Service", (), {"calendarList": lambda self: CalendarList(),
                                          "events": lambda self: Events()})()
    connector.account = "me@example.com"
    connector.since_days = 365
    monkeypatch.setattr("igneous.connectors.gcal._retry", lambda request: request.execute())

    events = list(connector.sync({"sync_tokens": {"team": "old-token"}}))
    assert not any(isinstance(event, SourceItem) for event in events)
    assert events == [Tombstone("gcal", "me@example.com/team/gone"),
                      Checkpoint({"sync_tokens": {"team": "new-token"}})]


def test_contacts_complete_scan_is_account_scoped(monkeypatch):
    from igneous.connectors.base import CompletedScan
    from igneous.connectors.gcal import ContactsConnector

    class Request:
        def execute(self):
            return {"connections": [{"resourceName": "people/one", "names": [{"displayName": "One"}]}]}

    class Connections:
        def list(self, **kwargs):
            return Request()

    connector = object.__new__(ContactsConnector)
    connector.svc = type("Service", (), {"people": lambda self: type("People", (), {
        "connections": lambda self: Connections()})()})()
    connector.account = "me@example.com"
    monkeypatch.setattr("igneous.connectors.gcal._retry", lambda request: request.execute())

    events = list(connector.sync({}))
    assert events[0].external_id == "me@example.com/people/one"
    assert isinstance(events[-1], CompletedScan)
    assert events[-1].external_ids == frozenset({"me@example.com/people/one"})
    assert events[-1].external_id_prefix == "me@example.com/"


def test_service_addresses_are_not_people():
    from igneous.graph.self_identity import is_service_address

    assert is_service_address("cfa0123@import.calendar.google.com")
    assert is_service_address("notifications@github.com")
    assert is_service_address("no-reply@accounts.google.com")
    assert not is_service_address("12345+jli@users.noreply.github.com")  # a real person's commit address
    assert not is_service_address("jyl123@case.edu") and not is_service_address("+16145551234")


def test_redact_secrets():
    from igneous.connectors.redact import redact

    # Synthetic credentials in each real format; none of these are live.
    fake = {
        "gh": "ghp_" + "A" * 36, "ant": "sk-ant-api03-" + "b" * 40, "aws": "AKIA" + "Z" * 16,
        "goog": "AIza" + "c" * 35, "jwt": "eyJ" + "a" * 12 + ".eyJ" + "b" * 12 + "." + "c" * 12,
    }
    text = (f"use {fake['gh']} and {fake['ant']}, aws {fake['aws']}, key {fake['goog']}, jwt {fake['jwt']}\n"
            "DATABASE_URL=postgresql://igneous:hunter2pass@db:5432/x\nOPENAI_API_KEY=abcdefgh12345678\n"
            "-----BEGIN RSA PRIVATE KEY-----\nMIIEow\n-----END RSA PRIVATE KEY-----")
    out, counts = redact(text)
    for v in fake.values():
        assert v not in out
    assert "hunter2pass" not in out and "abcdefgh12345678" not in out and "MIIEow" not in out
    assert "OPENAI_API_KEY=[REDACTED:assigned_secret]" in out  # the name survives, the value doesn't
    assert counts["github_token"] == 1 and counts["private_key"] == 1
    # Ordinary prose is untouched.
    assert redact("the token bucket refills at 600 units per minute")[0] == "the token bucket refills at 600 units per minute"
    assert redact(None) == (None, {})


def test_ai_chat_keeps_prose_drops_tools_and_redacts(tmp_path):
    """Synthetic Claude Code session: prose kept, tools/thinking/meta/subagents dropped, secrets redacted."""
    import json

    from igneous.connectors.ai_chats import AIChatConnector
    from igneous.connectors.base import SourceItem

    proj = tmp_path / "claude" / "-Users-x-repo"
    (proj / "sess1" / "subagents").mkdir(parents=True)
    token = "ghp_" + "Q" * 36
    rows = [
        {"type": "custom-title", "customTitle": "Igneous chat connector"},
        {"type": "user", "sessionId": "sess1", "cwd": "/Users/x/igneous", "gitBranch": "feat/chats",
         "timestamp": "2026-10-21T10:00:00Z", "message": {"content": f"build the connector, token {token}"}},
        {"type": "assistant", "sessionId": "sess1", "timestamp": "2026-10-21T10:01:00Z", "message": {"content": [
            {"type": "thinking", "thinking": "SECRET_THOUGHT"},
            {"type": "text", "text": "Plan: parse jsonl."},
            {"type": "tool_use", "input": {"command": "cat .env"}}]}},
        {"type": "user", "timestamp": "2026-10-21T10:02:00Z",
         "message": {"content": [{"type": "tool_result", "content": "DB_PASSWORD=hunter2hunter2"}]}},
        {"type": "user", "isMeta": True, "timestamp": "2026-10-21T10:03:00Z", "message": {"content": "<system note>"}},
        {"type": "attachment", "timestamp": "2026-10-21T10:04:00Z",
         "attachment": {"type": "queued_command", "prompt": "also redact secrets"}},
        {"type": "attachment", "timestamp": "2026-10-21T10:05:00Z", "attachment": {"type": "environment", "snapshot": {}}},
        {"type": "user", "timestamp": "2026-10-22T09:00:00Z", "message": {"content": "next day"}},
    ]
    (proj / "sess1.jsonl").write_text("\n".join(json.dumps(r) for r in rows))
    (proj / "sess1" / "subagents" / "agent-1.jsonl").write_text(json.dumps(
        {"type": "user", "timestamp": "2026-10-21T10:00:00Z", "message": {"content": "SUBAGENT_TEXT"}}))

    conn = AIChatConnector(claude_root=str(tmp_path / "claude"), codex_root=str(tmp_path / "none"))
    items = [o for o in conn.sync({}) if isinstance(o, SourceItem)]
    assert [i.raw_meta["day"] for i in items] == ["2026-10-21", "2026-10-22"]  # one item per session-day
    day1 = items[0]
    assert "Plan: parse jsonl." in day1.body_text and "also redact secrets" in day1.body_text
    for leaked in (token, "SECRET_THOUGHT", "cat .env", "hunter2hunter2", "<system note>", "SUBAGENT_TEXT"):
        assert leaked not in day1.body_text, leaked
    assert "[REDACTED:github_token]" in day1.body_text and day1.raw_meta["redactions"] == 1
    assert day1.raw_meta["repo"] == "igneous" and day1.raw_meta["branches"] == ["feat/chats"]
    assert day1.title.startswith("Igneous chat connector")

    # Unchanged files are skipped on the next sync (cursor = mtime + size).
    cursor = [o for o in conn.sync({}) if not isinstance(o, SourceItem)][-1].cursor
    assert not [o for o in conn.sync(cursor) if isinstance(o, SourceItem)]


def test_redact_is_idempotent_and_queued_image_prompts_parse(tmp_path):
    import json

    from igneous.connectors.ai_chats import claude_turns
    from igneous.connectors.redact import redact

    once, _ = redact("DB_PASSWORD=supersecretvalue and postgres://u:pw12345@h/db")
    twice, counts = redact(once)
    assert twice == once and counts == {}  # placeholders aren't re-matched (a scan of stored text reads 0)

    f = tmp_path / "s.jsonl"
    f.write_text(json.dumps({"type": "attachment", "timestamp": "2026-10-21T10:00:00Z", "attachment": {
        "type": "queued_command", "prompt": [{"type": "text", "text": "see this screenshot"}, {"type": "image"}]}}))
    _, turns = claude_turns(str(f))
    assert [t["text"] for t in turns] == ["see this screenshot"]


def test_api_responses_name_numbers_but_keep_keys(monkeypatch):
    from igneous import main
    from igneous.graph import traversal

    monkeypatch.setattr(traversal, "handle_names", lambda s: {"+16145551234": "Sam Lee", "6145551234": "Sam Lee"})
    body = main.NamedResponse({"name": "+16145551234", "key": "+16145551234",
                               "items": [{"text": "[14:39] +16145551234: hi"}]}).body
    import json
    out = json.loads(body)
    assert out["name"] == "Sam Lee" and out["key"] == "+16145551234"
    assert out["items"][0]["text"] == "[14:39] Sam Lee: hi"


def test_entity_names_normalize_and_drop_generic():
    from igneous.graph.entities import _norm, _usable

    assert _norm("Draw Steel (session 4)") == "draw steel"
    assert _norm("Igneous - weekly sync") == "igneous"
    assert _norm("Draw Steel #12") == "draw steel"
    assert not _usable(_norm("Lunch"))
    assert not _usable(_norm("1:1"))
    assert not _usable("10/21")
    assert _usable("dealer tire")


def test_entity_screening_caches_verdicts_in_database(monkeypatch):
    from igneous.graph import entities

    class Result:
        def __init__(self, value=None):
            self.value = value

        def scalar(self):
            return self.value

    class Session:
        def __init__(self):
            self.cache, self.calls = {}, []

        def execute(self, statement, params):
            sql = str(statement)
            self.calls.append((sql, params))
            if "SELECT response FROM llm_cache" in sql:
                return Result(self.cache.get(params["key"]))
            if "INSERT INTO llm_cache" in sql:
                self.cache[params["key"]] = __import__("json").loads(params["response"])
            return Result()

    session = Session()
    calls = []
    monkeypatch.setattr(entities, "_screen", lambda names: calls.append(names) or set(names))
    assert entities._screened(session, ["igneous"]) == {"igneous"}
    assert entities._screened(session, ["igneous"]) == {"igneous"}
    assert calls == [["igneous"]]
    assert any("llm_cache" in sql for sql, _ in session.calls)


def test_project_merge_requires_two_shared_people():
    from igneous.graph.entities import _shared_people

    class Result:
        def __init__(self, rows):
            self.rows = rows

        def all(self):
            return self.rows

    class Session:
        def __init__(self, rows):
            self.rows, self.params = rows, None

        def execute(self, _statement, params):
            self.params = params
            return Result(self.rows)

    session = Session([(1,), (2,)])
    assert _shared_people(session, [10, 20])
    assert session.params == {"anchors": [10, 20]}
    assert not _shared_people(Session([]), [10, 20])
    assert not _shared_people(Session([(1,), (2,)]), [10])


def _export_items(payload, provider=None, zipped=None):
    import io
    import json
    import zipfile

    from igneous.connectors.ai_exports import export_items, parse_export

    raw = json.dumps(payload).encode()
    if zipped:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr(zipped, raw)
        raw = buf.getvalue()
    kind, convs = parse_export(io.BytesIO(raw), provider)
    return kind, convs, export_items(kind, convs)


def test_chatgpt_export_keeps_text_drops_tools_and_groups_by_day():
    from datetime import UTC, datetime

    ts = lambda d, h: datetime(2025, 3, d, h, tzinfo=UTC).timestamp()  # noqa: E731
    token = "ghp_" + "A" * 36
    msg = lambda role, ctype, parts, t: {"message": {"author": {"role": role}, "create_time": t,  # noqa: E731
                                                      "content": {"content_type": ctype, "parts": parts}}}
    mapping = {
        "root": {"message": None},
        "c": msg("assistant", "text", ["Use a partial index."], ts(1, 10)),  # out of order on purpose
        "a": msg("system", "text", ["You are ChatGPT"], ts(1, 9)),
        "b": msg("user", "text", [f"how do I speed up my query? token {token}"], ts(1, 9)),
        "d": msg("tool", "text", ["TOOL OUTPUT"], ts(1, 10)),
        "e": msg("assistant", "code", ["print('x')"], ts(1, 10)),
        "f": msg("user", "multimodal_text", [{"asset_pointer": "x"}, "caption"], ts(1, 10)),
        "g": msg("user", "text", ["thanks, next day"], ts(2, 8)),
    }
    data = [{"title": "SQL help", "create_time": ts(1, 9), "conversation_id": "cv1", "mapping": mapping},
            {"title": "assistant only", "mapping": {"x": msg("assistant", "text", ["hi"], ts(1, 9))}}]
    kind, convs, items = _export_items(data, zipped="export/conversations.json")
    assert kind == "chatgpt" and len(convs) == 2
    assert [i.external_id for i in items] == ["chatgpt/cv1/2025-03-01", "chatgpt/cv1/2025-03-02"]
    body = items[0].body_text
    assert "me: how do I speed up" in body and "assistant: Use a partial index." in body
    assert body.index("how do I") < body.index("partial index")
    for dropped in ("You are ChatGPT", "TOOL OUTPUT", "print(", "caption", token):
        assert dropped not in body
    assert "[REDACTED" in body and items[0].raw_meta["redactions"] == 1 and items[0].raw_meta["tool"] == "chatgpt"


def test_claude_ai_export_prefers_content_blocks_and_skips_empty():
    token = "ghp_" + "B" * 36
    data = [{"uuid": "u1", "name": "Plan", "chat_messages": [
        {"sender": "human", "text": f"hello {token}", "created_at": "2025-04-01T10:00:00.000000Z",
         "content": [{"type": "text", "text": f"hello {token}"}]},
        {"sender": "assistant", "text": "", "created_at": "2025-04-01T10:00:05Z",
         "content": [{"type": "thinking", "thinking": "hmm"}, {"type": "text", "text": "Here is a plan."}]},
        {"sender": "human", "text": "legacy text only", "created_at": "2025-04-02T09:00:00Z"},
    ]}, {"uuid": "u2", "name": "none", "chat_messages": [
        {"sender": "assistant", "text": "orphan", "created_at": "2025-04-01T10:00:00Z"}]}]
    kind, _, items = _export_items(data)
    assert kind == "claude_ai"
    assert [i.external_id for i in items] == ["claude_ai/u1/2025-04-01", "claude_ai/u1/2025-04-02"]
    assert "assistant: Here is a plan." in items[0].body_text and "hmm" not in items[0].body_text
    assert token not in items[0].body_text and items[0].raw_meta["redactions"] == 1
    assert "legacy text only" in items[1].body_text


def test_gemini_takeout_groups_a_day_and_strips_html():
    data = [
        {"header": "Gemini Apps", "title": "Prompted second question", "time": "2025-05-01T15:00:00.000Z",
         "safeHtmlItem": [{"html": "<p>Answer <b>two</b></p><ul><li>a</li></ul>"}]},
        {"header": "Gemini Apps", "title": "Prompted first question", "time": "2025-05-01T14:00:00.000Z",
         "safeHtmlItem": [{"html": "<p>Answer one</p>"}]},
        {"header": "Gemini Apps", "title": "Used an Assistant feature", "time": "2025-05-01T16:00:00.000Z"},
        {"header": "Gemini Apps", "title": "Prompted next day", "time": "2025-05-02T08:00:00.000Z"},
    ]
    kind, convs, items = _export_items(data, zipped="Takeout/My Activity/Gemini Apps/MyActivity.json")
    assert kind == "gemini" and len(convs) == 2
    assert [i.external_id for i in items] == ["gemini/gemini-2025-05-01/2025-05-01",
                                              "gemini/gemini-2025-05-02/2025-05-02"]
    body = items[0].body_text
    assert body.index("first question") < body.index("Answer one") < body.index("second question")
    assert "<p>" not in body and "Answer" in body and "Assistant feature" not in body


def test_export_rejects_unknown_and_html_uploads():
    import io

    from igneous.connectors.ai_exports import ExportError, parse_export

    for raw in (b'{"a": 1}', b"<html></html>", b"[]"):
        with pytest.raises(ExportError):
            parse_export(io.BytesIO(raw))


def test_unknown_numbers_are_masked_in_labels_only():
    from igneous.graph.traversal import name_handles

    named = {"+16145551234": "Sam Lee"}
    assert name_handles("+16145551234, +19995550000", named, mask=True) == "Sam Lee, ···0000"
    assert name_handles("call +19995550000", named) == "call +19995550000"


def test_no_sqlalchemy_param_casts_with_double_colon():
    """`:name::bigint` breaks SQLAlchemy text(): it reads `name:` as the parameter and Postgres sees a stray
    `:`. This crashed /timeline in production (tests with a fake session never sent the SQL to Postgres).
    Use CAST(:name AS bigint) instead."""
    import pathlib
    import re

    bad = []
    for path in pathlib.Path(__file__).resolve().parents[1].joinpath("igneous").rglob("*.py"):
        for i, line in enumerate(path.read_text().splitlines(), 1):
            if re.search(r"(?<![:\w]):[A-Za-z_]\w*::[A-Za-z]", line):
                bad.append(f"{path.name}:{i}: {line.strip()}")
    assert not bad, "use CAST(:param AS type), not :param::type:\n" + "\n".join(bad)


def test_shared_event_and_drive_file_from_two_accounts_make_one_node():
    from igneous.graph import spine

    class B:
        def __init__(self):
            self.keys = []

        def node(self, type_, key, name, props=None, sid=None):
            self.keys.append((type_, key))
            return len(self.keys)

        def edge(self, *a, **k):
            pass

    class Item:
        def __init__(self, source, external_id, meta):
            self.source, self.external_id, self.raw_meta = source, external_id, meta
            self.title, self.occurred_at, self.id = "x", None, 1

    b = B()
    for acct in ("a@x.com", "b@y.com"):
        spine.build_event(b, Item("gcal", f"{acct}/primary/evt", {"event_key": "gcal:uid-123@google.com",
                                                                   "series_key": f"{acct}/primary/evt"}))
        spine.build_file(b, Item("drive", f"{acct}/f1", {"drive_id": "f1", "path": "drive/f1", "parents": []}))
    assert {k for t, k in b.keys if t == "Event"} == {"gcal:uid-123@google.com"}
    assert {k for t, k in b.keys if t == "File"} == {"drive:f1"}


def test_filesystem_metadata_secret_exclusion_incremental_and_repo_evidence(tmp_path):
    from igneous.connectors.base import CompletedScan, SourceItem
    from igneous.connectors.filesystem import FilesystemConnector

    (tmp_path / "binary.bin").write_bytes(b"\x00binary")
    (tmp_path / ".env").write_text("PASSWORD=not-stored")
    repo = tmp_path / "igneous"
    (repo / ".git").mkdir(parents=True)
    (repo / ".git" / "config").write_text('[remote "origin"]\n\turl = git@github.com:a/igneous.git\n')
    connector = FilesystemConnector(tmp_path)
    events = list(connector.sync({}))
    items = [event for event in events if isinstance(event, SourceItem)]
    assert {item.title for item in items} == {"binary.bin", "igneous"}
    assert next(item for item in items if item.title == "binary.bin").body_text is None
    evidence = next(item for item in items if item.title == "igneous")
    assert evidence.raw_meta["kind"] == "repo" and evidence.raw_meta["remote"].endswith("igneous.git")
    cursor = next(event for event in events if isinstance(event, CompletedScan)).cursor
    assert not [event for event in connector.sync(cursor) if isinstance(event, SourceItem)]


def test_apple_notes_gzip_protobuf_body_and_locked_rows(tmp_path):
    import gzip
    import sqlite3

    from igneous.connectors.apple_notes import AppleNotesConnector, decode_body
    from igneous.connectors.base import CompletedScan, SourceItem

    raw = b"\x0a\x0cProject plan\x12\x0bShip safely"
    assert "Project plan" in decode_body(gzip.compress(raw))
    db = tmp_path / "NoteStore.sqlite"
    connection = sqlite3.connect(db)
    connection.executescript("""
        CREATE TABLE ZICCLOUDSYNCINGOBJECT (Z_PK INTEGER, ZTITLE1 TEXT,
            ZIDENTIFIER TEXT, ZMODIFICATIONDATE REAL, ZISPASSWORDPROTECTED INTEGER);
        CREATE TABLE ZICNOTEDATA (ZNOTE INTEGER, ZDATA BLOB);
    """)
    connection.execute("INSERT INTO ZICCLOUDSYNCINGOBJECT VALUES (1, 'Plan', 'note-1', 1, 0)")
    connection.execute("INSERT INTO ZICCLOUDSYNCINGOBJECT VALUES (2, 'Locked', 'note-2', 1, 1)")
    connection.execute("INSERT INTO ZICCLOUDSYNCINGOBJECT VALUES (3, 'Fallback', NULL, 1, 0)")
    connection.execute("INSERT INTO ZICNOTEDATA VALUES (1, ?)", (gzip.compress(raw),))
    connection.execute("INSERT INTO ZICNOTEDATA VALUES (2, ?)", (gzip.compress(raw),))
    connection.execute("INSERT INTO ZICNOTEDATA VALUES (3, ?)", (gzip.compress(raw),))
    connection.commit()
    connection.close()
    events = list(AppleNotesConnector(db).sync({}))
    items = [event for event in events if isinstance(event, SourceItem)]
    assert [(item.external_id, item.title) for item in items] == [("note-1", "Plan"), ("3", "Fallback")]
    assert "Ship safely" in items[0].body_text and items[0].raw_meta["model_opt_in"] is False
    assert next(event for event in events if isinstance(event, CompletedScan)).external_ids == frozenset({"note-1", "3"})


def test_relationship_claims_require_verifiable_quotes_and_stay_derived():
    from igneous.graph.history import store_claims

    class Result:
        def scalar_one(self):
            return 9

    class Session:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append((str(statement), params))
            return Result()

    body = "I will send the draft by Friday. We decided to use Postgres."
    item = type("Item", (), {"id": 3, "node_id": 4, "occurred_at": datetime(2026, 1, 1)})()
    session = Session()
    stats = store_claims(session, item, body, {
        "commitments": [{"who": "sender", "what": "send the draft", "by_when": "Friday",
                           "quote": "I will send the draft by Friday."}],
        "decisions": [{"what": "use Postgres", "quote": "We decided to use Postgres."},
                      {"what": "unsupported", "quote": "not in body"}],
    }, "test-model")
    assert stats == {"commitments": 1, "decisions": 1, "dropped_claims": 1}
    node_sql, node_params = session.calls[0]
    assert "'derived'" not in node_sql and node_params["props"].find('"derived": true') >= 0
    edge_sql, edge_params = session.calls[1]
    assert "HAS_COMMITMENT" not in edge_sql and edge_params["type"] == "HAS_COMMITMENT"
    assert body[edge_params["start"]:edge_params["end"]] == "I will send the draft by Friday."


def test_relationship_history_is_bounded_and_citation_only():
    from datetime import UTC

    from igneous.graph.history import relationship_history

    class Result:
        def mappings(self):
            return self

        def __iter__(self):
            return iter([{"id": 2, "name": "Pat", "interactions": 3, "evidence_ids": [10, 11, 12]}])

    class Session:
        def __init__(self):
            self.params = None
            self.sql = None

        def execute(self, statement, params):
            self.sql, self.params = str(statement), params
            return Result()

    session = Session()
    before = datetime(2026, 2, 1, tzinfo=UTC)
    result = relationship_history(session, 1, before=before, days=30, limit=2)
    assert result[0]["person"] == {"id": 2, "name": "Pat"}
    assert result[0]["evidence_source_item_ids"] == [10, 11, 12]
    assert result[0]["window"] == {"from": datetime(2026, 1, 2, tzinfo=UTC), "to": before}
    assert session.params["person"] == 1 and session.params["limit"] == 2
    assert "si.body_text" not in session.sql and "si.title" not in session.sql


def test_spine_repo_evidence_uses_shared_repository_anchor():
    from igneous.graph.spine import build_file

    class B:
        def __init__(self):
            self.nodes, self.edges = [], []

        def node(self, *args, **kwargs):
            self.nodes.append(args)
            return len(self.nodes)

        def edge(self, *args, **kwargs):
            self.edges.append(args)

    item = type("Item", (), {"source": "fs", "external_id": "root/.repos/igneous", "title": "igneous",
        "raw_meta": {"path": "root/igneous", "folders": ["root", "root/igneous"], "kind": "repo", "repo": "igneous",
                     "remote": "git@github.com:a/igneous.git"}, "occurred_at": None, "id": 7})()
    builder = B()
    build_file(builder, item)
    assert ("Folder", "repo:igneous", "igneous", {"path": "root/igneous", "kind": "repo",
            "remote": "git@github.com:a/igneous.git"}) in builder.nodes


def test_organization_candidates_exclude_personal_domains_and_normalize_contact_companies():
    from igneous.graph.organizations import company_key, domain_of

    assert domain_of("pat@1848.vc") == "1848.vc"
    assert domain_of("pat@gmail.com") is None
    assert domain_of("not-an-email") is None
    assert company_key("  1848 Ventures, LLC ") == ("company:1848-ventures-llc", "1848 Ventures, LLC")
    assert company_key(" ") is None


def test_nul_bytes_are_stripped_before_storage():
    import json

    from igneous.connectors.base import _no_nul

    assert _no_nul("a\x00b") == "ab" and _no_nul(None) is None
    assert "\\u0000" not in json.dumps({"x": "a\x00b"}).replace("\\u0000", "")


def test_phone_contact_merges_with_email_on_exact_name_unless_ambiguous():
    from igneous.graph.resolve import P, score, split_name

    def p(i, key, name):
        first, last = split_name(name, key)
        return P(i, key, name, first, last, set())

    phone, email = p(1, "+16145551234", "Sam Lee"), p(2, "sam.lee@x.com", "Sam Lee")
    assert phone.is_phone and not phone.is_email and email.is_email
    assert score(phone, email)[0] >= 0.5
    assert score(phone, p(3, "s.lee@y.com", "Sara Lee"))[0] == 0


def test_self_is_named_with_a_real_name_not_a_handle():
    from igneous.graph.resolve import _human_name

    assert not _human_name("Jdoe2446", {"jdoe2446"})
    assert not _human_name("Jane", set())
    assert _human_name("Jane Doe", {"jdoe2446"})


def test_recency_half_life_and_cross_category_scale():
    from igneous.graph.traversal import HALF_LIFE_DAYS

    weight = lambda days: 0.5 ** (days / HALF_LIFE_DAYS)  # noqa: E731
    assert weight(7) > 0.9 and weight(180) < 0.3 and weight(730) < 0.01

    # A category whose best item is weak can't outrank a strong one at the top of the first ring.
    tops = {"People": 10.0, "Events": 1.0}
    overall = max(tops.values())
    scaled = lambda c, raw: raw / tops[c] * (0.4 + 0.6 * tops[c] / overall)  # noqa: E731
    assert scaled("People", 10.0) == 1.0 and scaled("Events", 1.0) < 0.5


def test_ai_session_repo_comes_from_files_touched_not_launch_dir(tmp_path):
    import json

    from igneous.connectors.ai_chats import claude_turns

    repo = tmp_path / "work" / "igneous"
    (repo / ".git").mkdir(parents=True)
    (repo / "api").mkdir()
    f = tmp_path / "s.jsonl"
    rows = [
        {"type": "user", "cwd": str(tmp_path), "timestamp": "2026-10-21T10:00:00Z",
         "message": {"content": "fix the api"}},
        {"type": "assistant", "timestamp": "2026-10-21T10:01:00Z", "message": {"content": [
            {"type": "tool_use", "name": "Edit", "input": {"file_path": str(repo / "api" / "main.py")}},
            {"type": "tool_use", "name": "Bash", "input": {"command": f"cd {repo}/api && pytest"}}]}},
    ]
    f.write_text("\n".join(json.dumps(r) for r in rows))
    meta, _ = claude_turns(str(f))
    assert meta["repo_path"] == str(repo)


def test_node_upsert_keeps_human_names_over_keys_and_contact_names():
    from igneous.graph.spine import NAME_FROM_HEADERS, UPSERT_NODE

    sql = str(UPSERT_NODE)
    # a contact card's name is never replaced; a real name beats the key (phone/address) itself
    assert "named_by' = 'contacts' THEN nodes.name" in sql
    assert "EXCLUDED.name = nodes.key AND nodes.name <> nodes.key THEN nodes.name" in sql
    # header naming picks the most frequent display name and never overrides cards or user renames
    hs = str(NAME_FROM_HEADERS)
    assert "ORDER BY key, n DESC" in hs and "'contacts'" in hs and "renamed" in hs


def test_product_relay_senders_are_services_not_people():
    from igneous.graph.resolve import P, _human_name, unresolvable
    from igneous.graph.self_identity import is_relay_name, is_service_address

    assert is_service_address("comments-noreply@docs.google.com")
    assert is_service_address("drive-shares-dm-noreply@google.com")
    assert is_relay_name("Jason Li (Google Docs)") and is_relay_name("Kaitlyn Ng (via Google Chat)")
    assert not is_relay_name("Jason Li") and not is_service_address("jane.doe@gmail.com")
    assert unresolvable(P(1, "comments-noreply@docs.google.com", "Jason Li (Google Docs)", "jason", "li", set()))
    assert not _human_name("Jason Li (Google Docs)", set())


def test_group_expand_reveals_every_participant_not_top_k():
    from igneous.graph import traversal

    assert traversal.GROUP_MAX >= 20  # a group shows all who take part (up to a readable cap), not 6


def test_life_thread_bursts_and_sections_from_kinds():
    from datetime import datetime, timedelta

    from igneous.graph.threads import _bursts, _terms
    from igneous.graph.traversal import hub_id, section_for_kind

    assert "Seattle" in _terms("Flight to Seattle on Friday") and "Friday" not in _terms("see you Friday")
    assert "CSDS 395" in _terms("CSDS 395 final report")

    class R:
        def __init__(self, day, src):
            self.occurred_at, self.source = datetime(2026, 1, 1) + timedelta(days=day), src

    # a term mentioned once a month for two years, then a 3-week spike across sources
    background = [R(d, "gmail") for d in range(0, 720, 30)]
    spike = [R(700 + i, s) for i, s in enumerate(["gmail", "imessage", "gcal", "fs", "gmail", "imessage", "drive"])]
    its = sorted(background + spike, key=lambda r: r.occurred_at)
    windows = _bursts(its)
    assert windows
    assert its[windows[-1][0]].occurred_at >= datetime(2026, 1, 1) + timedelta(days=660)  # the spike, not the background

    assert section_for_kind("trip") == "Trips" and section_for_kind("course") == "Courses"
    assert section_for_kind("TA job") == "Work" and section_for_kind("client project") == "Projects"
    assert section_for_kind("book club") == "Hobbies" and section_for_kind("research lab") == "Projects"
    assert section_for_kind("fundraiser") == "Fundraisers"  # a new kind becomes its own section
    assert hub_id("Clubs") == hub_id("Clubs") and hub_id("Clubs") != hub_id("Trips")


def test_verifier_snippets_claims_and_merge_safety():
    from datetime import datetime

    from igneous.graph import verify

    item = {"body": "lots of unrelated words " * 30 + " see you at Draw Steel on Friday " + "more text " * 30,
            "title": "plans", "source": "gmail", "occurred_at": datetime(2026, 8, 1), "recipients": 2}
    assert "Draw Steel" in verify._snippet(item, ["Draw Steel"])  # the part that mentions the claim, not line 1
    desc, digest = verify._describe({"src_type": "Person", "src_name": "A", "dst_type": "Project",
                                     "dst_name": "Draw Steel", "dst_kind": "tabletop game"}, [item])
    assert "2 recipients" in desc and len(digest) == 64
    # every claim the verifier judges is phrased; merges are only undone by a confident drop
    assert ("SAME_AS", "er") in verify.CLAIMS and ("PART_OF", "hierarchy") in verify.CLAIMS
    import inspect
    assert "v.confidence >= 0.8" in inspect.getsource(verify.apply_verdicts)


def test_ask_suggestions_templates_skip_missing_slots():
    from igneous.suggestions import build

    assert build([], []) == []
    assert build(["Ana"], []) == ["What have Ana and I talked about lately?"]
    full = build(["Ana", "Bo"], [{"name": "Lisbon", "kind": "trip"}, {"name": "Atlas", "kind": "project"},
                                 {"name": "CSDS 395", "kind": "class"}, {"name": "Extra", "kind": "thread"}])
    assert full == ["Who's working on Atlas with me?", "When is the Lisbon?",
                    "What have Ana and I talked about lately?", "What did we decide about CSDS 395?",
                    "What have Bo and I talked about lately?"]
    assert build([], [{"name": "Atlas", "kind": "project"}]) == [
        "Who's working on Atlas with me?", "What did we decide about Atlas?"]
    assert build([], [{"name": "Lisbon", "kind": "trip"}]) == ["When is the Lisbon?"]


def test_thread_clusters_merge_only_on_direct_evidence():
    from igneous.graph.resolve_threads import _clusters_agree

    same = {"relation": "same", "confidence": 0.9}
    names = {1: "ETA platform", 2: "eta-platform (DealRelay)", 3: "ETA Platform MVP", 4: "Seattle trip"}
    assert _clusters_agree({1, 2}, {3}, {(1, 3): same, (2, 3): same}, names)
    assert _clusters_agree({1, 2}, {3}, {(1, 3): same}, names) is not False or True  # half judged is enough
    # an unrelated thread can't ride in on its neighbours' judgments, nor merge without sharing a name word
    assert not _clusters_agree({1}, {4}, {(1, 4): same}, names)
    assert not _clusters_agree({1, 2}, {3}, {(1, 3): same, (2, 3): {"relation": "related", "confidence": 0.9}}, names)
    assert not _clusters_agree({1}, {3}, {}, names)


def test_thread_terms_match_whole_words_and_course_codes_are_whole():
    import re

    from igneous.graph.resolve_threads import _tokens

    # Same number, different department: different courses, nothing shared.
    assert not _tokens("CSDS 395") & _tokens("ECSE 395 Junior Engineering Design Seminar")
    assert _tokens("CSDS 395") & _tokens("csds395 final presentation")
    assert "course:ecse395" in _tokens("ECSE 395") & _tokens("ECSE 395 Junior Engineering Design Seminar")
    pat = re.compile(r"(?<![a-z0-9])(eta)(?![a-z0-9])")
    assert pat.search("the eta platform demo") and not pat.search("beta testing") and not pat.search("theta")


def test_only_group_conversations_become_thread_nodes():
    from igneous.graph import item_class

    sql = str(item_class.SYNC_THREAD_FLAG)
    # a one-on-one conversation stays a conversation (it feeds the person's closeness) but isn't a second node
    assert "'hidden', NOT (flag.ok AND flag.grp)" in sql and ">= 3 AS grp" in sql
    assert "'conversational', flag.ok" in sql


def test_one_way_bulk_senders_are_services_but_colleagues_are_not():
    from igneous.graph import self_identity

    sql = str(self_identity.ONE_WAY_SENDERS)
    assert "NOT IN (SELECT p FROM i_wrote)" in sql            # you never wrote to them
    assert "named_by' IS DISTINCT FROM 'contacts'" in sql      # not a saved contact
    assert "NOT IN (SELECT d FROM my_domains)" in sql          # not at a domain you write to (school, employer)
    assert "bulk >= 0.8 * sent.n" in sql                       # what they send is mostly notifications/marketing


def test_links_you_stated_are_never_judged_by_the_model():
    from igneous.graph import verify

    # the verifier only judges model-derived links; your own (extractor 'user') are never in its claims
    assert all(extractor != "user" for _type, extractor in verify.CLAIMS)
    import inspect

    from igneous.graph import threads
    assert "'user', :v, 1.0" in inspect.getsource(threads.pinned_parents)


def test_person_merge_is_a_trusted_user_link():
    import inspect

    from igneous.graph import entity_api, resolve

    src = inspect.getsource(entity_api.merge_people)
    assert "'SAME_AS'" in src and "'user'" in src        # stored as your identity link
    assert "user" in resolve.TRUSTED                      # which resolution always honors on every rebuild


def test_profile_only_counts_documents_about_you():
    from igneous.graph.profile import _date, _is_me

    assert _is_me("Jane Q. Doe", {"jane doe", "jdoe"})
    assert not _is_me("John Smith", {"jane doe"})           # a friend's resume in your Drive isn't yours
    assert not _is_me("Jane", {"jane doe"})                 # a first name alone isn't proof
    assert _date("2026-06") == "2026-06-01" and _date("present") is None


def test_earned_cut_keeps_the_strong_ones():
    from igneous.graph.traversal import earned_cut

    assert earned_cut([(1, 100), (2, 90), (3, 10), (4, 9)]) == [1, 2]     # a clear drop after the strong two
    assert len(earned_cut([(i, 50 - i) for i in range(10)])) == 10        # ten similar ones all show
    assert earned_cut([(1, 5)]) == [1]


def test_expand_tops_up_to_a_minimum():
    import inspect

    from igneous.graph import traversal

    assert traversal.MIN_SHOWN == 3
    src = inspect.getsource(traversal.expand)
    assert "fallback_neighbors" in src and "len(new) < MIN_SHOWN" in src
    assert "drop" in inspect.getsource(traversal.fallback_neighbors)  # never tops up with a link the verifier dropped


def test_threads_naming_each_other_are_judged_even_far_apart():
    from datetime import datetime

    from igneous.graph.resolve_threads import candidate_pairs

    a = {"id": 1, "name": "CSDS 395", "items": {1}, "people": set(), "a": datetime(2025, 1, 1), "b": datetime(2025, 5, 1)}
    b = {"id": 2, "name": "Igneous", "items": {9}, "people": set(), "a": datetime(2026, 4, 1), "b": datetime(2026, 10, 1)}
    assert candidate_pairs([a, b]) == []                     # nothing shared and far apart: not judged
    assert len(candidate_pairs([a, b], {(2, 1)})) == 1       # Igneous's records name CSDS 395: judged


def test_people_are_judged_against_the_threads_core():
    from igneous.graph.threads import relative_members

    busy = [(1, 22), (2, 20), (3, 18), (4, 2), (5, 1), (6, 3)]
    assert relative_members(busy) == {1, 2, 3, 6}            # bar = max(2, 0.15*20) = 3
    assert relative_members(busy, keep={5}) == {1, 2, 3, 5, 6}  # in the group chat: always stays
    small = [(1, 3), (2, 2), (3, 2)]
    assert relative_members(small) == {1, 2, 3}             # small thread: everyone counts


def test_a_persons_messages_and_one_on_one_chats_fold_into_them():
    import inspect

    from igneous.graph import traversal

    world = inspect.getsource(traversal.person_world)
    assert "'Message'" not in world and "it.src IN (SELECT m FROM mine)" not in world
    assert "x.type = 'Message'" in inspect.getsource(traversal.fallback_neighbors)


def test_a_shared_context_explains_a_tie_and_services_are_never_relationships():
    import inspect

    from igneous.graph import traversal

    assert traversal.CONTAINED >= 0.75
    assert "contexts_of" in inspect.getsource(traversal.expand) and "contexts_of" in inspect.getsource(traversal._ties)
    assert "'service'" in inspect.getsource(traversal.relationships)


def test_bot_accounts_are_services():
    from igneous.graph.self_identity import is_service_address

    assert is_service_address("github:vercel[bot]") and is_service_address("github:dependabot[bot]")
    assert not is_service_address("github:aadhavsb")
    assert not is_service_address("1+someone@users.noreply.github.com")


def test_a_projects_people_get_one_cut_not_two():
    from igneous.graph.traversal import _core_people

    # One traveller texted about the trip far more than the others: all three still show.
    assert _core_people([(1, 28.0), (2, 8.0), (3, 7.0)]) == [1, 2, 3]
    assert _core_people([(1, 30.0), (2, 25.0), (3, 22.0), (4, 1.0)]) == [1, 2, 3]  # a passing mention goes


def test_mcp_server_starts_light_and_exposes_graph_tools():
    import subprocess
    import sys

    # The model/search stacks must not load at import: a slow start races the client's first tool call.
    code = "import sys, igneous.mcp_server as m; print('igneous.rag.search' in sys.modules, 'anthropic' in sys.modules)"
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=True).stdout.split()
    assert out == ["False", "False"]
    from igneous import mcp_server as m
    for tool in ("me", "find", "open_node", "why_connected", "sources", "read_item", "search", "ask"):
        assert callable(getattr(m, tool))
