"""Documents are chunked along the pages, sections, slides and sheets the parser recorded (#186)."""
from io import BytesIO
from types import SimpleNamespace

from igneous.rag import chunking

SECRET = "sk-ant-" + "a" * 40


def item(body, boundaries, source="drive", title="plan.pdf", **meta):
    return SimpleNamespace(source=source, title=title, body_text=body, occurred_at=None,
                           raw_meta={"boundaries": boundaries, **meta})


def pages(*texts):
    """Body and page boundaries exactly as parsers._parse_pdf builds them."""
    bounds, offset = [], 0
    for i, t in enumerate(texts, start=1):
        if i > 1:
            offset += 2
        bounds.append({"kind": "page", "index": i, "start": offset, "end": offset + len(t)})
        offset += len(t)
    return "\n\n".join(texts), bounds


def test_small_pages_pack_whole_and_each_chunk_says_where_it_is():
    body, bounds = pages("Budget overview", "Hiring plan", "Risks")
    assert chunking.chunks_for(item(body, bounds)) == [
        "plan.pdf, page 1 to page 3\nBudget overview\n\nHiring plan\n\nRisks"]


def test_a_page_never_shares_a_chunk_unless_packed_whole():
    long_page = "word " * 600   # 3000 characters: split on its own, with overlap
    body, bounds = pages("Intro", long_page.strip(), "Appendix")
    chunks = chunking.chunks_for(item(body, bounds))
    assert chunks[0] == "plan.pdf, page 1\nIntro"
    assert all(c.startswith("plan.pdf, page 2\n") for c in chunks[1:-1]) and len(chunks) >= 4
    assert chunks[-1] == "plan.pdf, page 3\nAppendix"


def test_docx_sections_run_from_heading_to_heading_and_tables_follow():
    from docx import Document

    from igneous.parsers import parse_document

    doc = Document()
    doc.add_paragraph("Draft for the team")
    doc.add_heading("Background", level=1)
    doc.add_paragraph("Why we started")
    doc.add_heading("Plan", level=1)
    doc.add_paragraph("Ship in May")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "owner", "Sam"
    data = BytesIO()
    doc.save(data)
    parsed = parse_document(data.getvalue(), "proposal.docx")
    units = chunking.sections(item(parsed.text, parsed.boundaries, title="proposal.docx"))
    assert units == [("", "Draft for the team"), ("Background", "Background\nWhy we started"),
                     ("Plan", "Plan\nShip in May"), ("table 1", "owner\tSam")]


def test_sheets_are_named():
    body = "a\t1\n\nb\t2"
    bounds = [{"kind": "sheet", "index": 1, "name": "Budget", "start": 0, "end": 3},
              {"kind": "sheet", "index": 2, "name": "Staff", "start": 5, "end": 8}]
    assert chunking.sections(item(body, bounds, title="q3.xlsx")) == [("sheet Budget", "a\t1"), ("sheet Staff", "b\t2")]


def test_untrusted_boundaries_fall_back_to_plain_splitting():
    body, bounds = pages("one", "two")
    # A local file imported before boundaries followed redaction: its offsets point into the unredacted text.
    old_local = item(body, bounds, source="fs", redactions={"anthropic_key": 1})
    assert chunking.chunks_for(old_local) == [body]
    assert chunking.chunks_for(item(body, [{"kind": "page", "index": 1, "start": 0, "end": 999}])) == [body]
    assert chunking.chunks_for(item(body, [])) == [body]


def test_redaction_moves_a_local_files_page_boundaries(tmp_path):
    """The fs connector redacts after parsing: every replacement shifts what follows it."""
    from pptx import Presentation

    from igneous.connectors.base import SourceItem
    from igneous.connectors.filesystem import FilesystemConnector

    deck = Presentation()
    for words in (f"Setup uses {SECRET} for now", "Second slide", "Third slide"):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.add_textbox(0, 0, 1, 1).text = words
    deck.save(tmp_path / "deck.pptx")
    (doc,) = [x for x in FilesystemConnector(tmp_path).sync({}) if isinstance(x, SourceItem)]
    assert SECRET not in doc.body_text and doc.raw_meta["redactions"] == {"anthropic_key": 1}
    assert [doc.body_text[b["start"]:b["end"]] for b in doc.raw_meta["boundaries"]] == [
        "Setup uses [REDACTED:anthropic_key] for now", "Second slide", "Third slide"]
    assert chunking.chunks_for(doc) == [
        "deck.pptx, slide 1 to slide 3\nSetup uses [REDACTED:anthropic_key] for now\n\nSecond slide\n\nThird slide"]


def test_documents_chunked_the_old_way_are_rechunked_once(test_db):
    import json
    from datetime import UTC, datetime

    from sqlalchemy import text

    from igneous.db import SessionLocal
    from igneous.pipeline import advance, watermark

    body, bounds = pages("First page", "Second page")
    with SessionLocal() as s:
        s.execute(text("TRUNCATE source_items, nodes, chunks, pipeline_state RESTART IDENTITY CASCADE"))
        ids = {}
        for source, ext, raw, chunk in (("drive", "doc", {"boundaries": bounds}, body),
                                        ("gmail", "mail", {}, "Subject: Hi\nhello")):
            ids[source] = sid = s.execute(text("""
                INSERT INTO source_items (source, external_id, content_hash, title, body_text, raw_meta)
                VALUES (:s, :e, :e, :t, :b, CAST(:m AS jsonb)) RETURNING id"""),
                {"s": source, "e": ext, "t": "plan.pdf" if source == "drive" else "Hi",
                 "b": body if source == "drive" else "hello", "m": json.dumps(raw)}).scalar_one()
            nid = s.execute(text("INSERT INTO nodes (type, key, name, source_item_id) VALUES "
                                 "(:t, :k, :k, :i) RETURNING id"),
                            {"t": "File" if source == "drive" else "Message", "k": ext, "i": sid}).scalar_one()
            s.execute(text("INSERT INTO chunks (source_item_id, node_id, ord, text, meta) VALUES (:i, :n, 0, :c, '{}')"),
                      {"i": sid, "n": nid, "c": chunk})  # chunked before layouts existed
        watermark(s, chunking.STAGE)
        advance(s, chunking.STAGE, datetime.now(UTC))  # both are behind the watermark
        s.commit()
        assert chunking.run(s) == 1   # only the document
        doc = s.execute(text("SELECT text, meta->>'layout' FROM chunks WHERE source_item_id = :i"),
                        {"i": ids["drive"]}).all()
        assert doc == [("plan.pdf, page 1 to page 2\nFirst page\n\nSecond page", str(chunking.LAYOUT))]
        assert s.execute(text("SELECT text FROM chunks WHERE source_item_id = :i"),
                         {"i": ids["gmail"]}).scalar() == "Subject: Hi\nhello"  # untouched, embedding kept
        assert chunking.run(s) == 0   # once
        s.execute(text("TRUNCATE source_items, nodes, chunks, pipeline_state RESTART IDENTITY CASCADE"))
        s.commit()
